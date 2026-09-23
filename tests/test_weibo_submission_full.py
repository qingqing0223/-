from __future__ import annotations

from datetime import datetime
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook

from monitor.weibo_submission_full import export_weibo_submission


class WeiboSubmissionFullTests(unittest.TestCase):

    def test_five_tables_manifest_and_excel(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)

            raw_root = root / "raw"
            raw_root.mkdir()

            content_path = raw_root / "search_contents.jsonl"
            comment_path = raw_root / "comments.jsonl"
            profile_path = raw_root / "creator_account_info.jsonl"

            content = {
                "note_id": "5345000000000001",
                "content": (
                    "2026\u5e74\u6c11\u65cf\u56e2\u7ed3\u8fdb\u6b65"
                    "\u5ba3\u4f20\u5468\u76f8\u5173\u5ba3\u4f20\u5185\u5bb9"
                ),
                "create_date_time": "2026-09-20T10:00:00+08:00",
                "liked_count": 10,
                "comments_count": 2,
                "shared_count": 3,
                "note_url": (
                    "https://m.weibo.cn/detail/5345000000000001"
                ),
                "user_id": "5055720243",
                "nickname": "\u4e2d\u56fd\u6c11\u65cf\u62a5",
                "ip_location": "\u5317\u4eac",
                "source_keyword": (
                    "2026\u5e74\u6c11\u65cf\u56e2\u7ed3\u8fdb\u6b65\u5ba3\u4f20\u5468"
                ),
                "original_or_repost": "\u8f6c\u8f7d",
            }

            comment = {
                "comment_id": "6000000000000001",
                "note_id": "5345000000000001",
                "content": "\u8f6c\u53d1\u4e86\u89e3",
                "create_date_time": "2026-09-20T10:10:00+08:00",
                "user_id": "9000000001",
                "nickname": "\u7528\u6237\u7532",
                "ip_location": "\u4e0a\u6d77",
                "sub_comment_count": 1,
                "comment_like_count": 4,
                "parent_comment_id": "",
                "root_comment_id": "6000000000000001",
            }

            profile = {
                "account_id": "5055720243",
                "account_name": "\u4e2d\u56fd\u6c11\u65cf\u62a5",
                "profile_url": "",
                "followers": "153.2\u4e07",
                "following": "374",
                "region": "",
            }

            content_path.write_text(
                json.dumps(content, ensure_ascii=False) + "\n",
                encoding="utf-8",
            )

            comment_path.write_text(
                json.dumps(comment, ensure_ascii=False) + "\n",
                encoding="utf-8",
            )

            profile_path.write_text(
                json.dumps(profile, ensure_ascii=False) + "\n",
                encoding="utf-8",
            )

            config_dir = root / "config"
            config_dir.mkdir()

            (config_dir / "key_accounts.v3.catalog.json").write_text(
                json.dumps(
                    {
                        "categories": [
                            {
                                "account_type": "\u7edf\u6218\u7cfb\u7edf\u5a92\u4f53",
                                "accounts": ["\u4e2d\u56fd\u6c11\u65cf\u62a5"],
                            }
                        ]
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            (config_dir / "weibo_scope_overrides.json").write_text(
                json.dumps(
                    {
                        "force_valid_content_ids": [],
                        "exclude_content_ids": {},
                    }
                ),
                encoding="utf-8",
            )

            cfg = {
                "monitoring_start_time": "2026-09-16T00:00:00+08:00",
                "key_account_catalog": "config/key_accounts.v3.catalog.json",
                "keywords": [
                    "2026\u5e74\u6c11\u65cf\u56e2\u7ed3\u8fdb\u6b65\u5ba3\u4f20\u5468"
                ],
            }

            result = export_weibo_submission(
                [content_path, comment_path],
                cfg,
                root,
                "wb01",
                raw_root=raw_root,
                now=datetime.fromisoformat(
                    "2026-09-21T09:05:00+08:00"
                ),
            )

            out = Path(result["output_dir"])

            def rows(name):
                return [
                    json.loads(line)
                    for line in (
                        out / name
                    ).read_text(
                        encoding="utf-8"
                    ).splitlines()
                    if line.strip()
                ]

            t1 = rows("table1_content.jsonl")
            t2 = rows("table2_comments.jsonl")
            t3 = rows("table3_post_interactions.jsonl")
            t4 = rows("table4_comment_interactions.jsonl")
            t5 = rows("table5_accounts.jsonl")

            self.assertEqual(
                [len(t1), len(t2), len(t3), len(t4), len(t5)],
                [1, 1, 1, 1, 1],
            )

            self.assertEqual(t1[0]["author_id"], "5055720243")
            self.assertEqual(t1[0]["original_or_repost"], "\u8f6c\u8f7d")

            self.assertEqual(t2[0]["comment_user_id"], "9000000001")
            self.assertEqual(t2[0]["comment_user_name"], "\u7528\u6237\u7532")

            self.assertEqual(t3[0]["like_count"], 10)
            self.assertEqual(t3[0]["comment_count"], 2)
            self.assertEqual(t3[0]["repost_count"], 3)

            self.assertEqual(t3[0]["view_count"], "")
            self.assertEqual(t3[0]["share_count"], "")
            self.assertEqual(t3[0]["favorite_count"], "")

            self.assertEqual(t4[0]["comment_reply_count"], 1)
            self.assertEqual(t4[0]["comment_like_count"], 4)

            self.assertEqual(t5[0]["followers"], 1532000)
            self.assertEqual(t5[0]["following"], 374)
            self.assertEqual(
                t5[0]["account_type"],
                "\u7edf\u6218\u7cfb\u7edf\u5a92\u4f53",
            )
            self.assertEqual(t5[0]["total_interactions"], 15)

            workbook = load_workbook(
                result["excel"],
                read_only=True,
            )

            self.assertEqual(
                workbook.sheetnames,
                [
                    "\u8bf4\u660e\u4e0e\u7edf\u8ba1",
                    "\u88681-\u53d1\u5e03\u5185\u5bb9",
                    "\u88682-\u8bc4\u8bba",
                    "\u88683-\u53d1\u5e03\u4e92\u52a8",
                    "\u88684-\u8bc4\u8bba\u4e92\u52a8",
                    "\u88685-\u5e10\u53f7\u4fe1\u606f",
                ],
            )

            workbook.close()

    def test_same_hour_snapshot_is_replaced(self):
        self.assertEqual(
            __import__(
                "monitor.weibo_submission_full",
                fromlist=["_hour"],
            )._hour("2026-09-21T09:59:59+08:00"),
            "2026-09-21T09:00:00+08:00",
        )


if __name__ == "__main__":
    unittest.main()