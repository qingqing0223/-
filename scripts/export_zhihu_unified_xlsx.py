from __future__ import annotations

"""Export Zhihu TechDesign V3 tables 1-5 into the unified workbook template."""

import argparse
import json
from pathlib import Path
import shutil

from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE = REPO_ROOT.parent / "数据" / "舆情监测统一数据模板v1.xlsx"
DEFAULT_OUTPUT = REPO_ROOT.parent / "数据" / "9-22" / "知乎舆情监测统一数据_2026-09-22.xlsx"

TABLE_FILES = {
    "表1-发布内容": "table1_content.jsonl",
    "表2-评论基础": "table2_comments.jsonl",
    "表3-发布互动": "table3_content_engagement.jsonl",
    "表4-评论互动": "table4_comment_engagement.jsonl",
    "表5-账号信息": "table5_accounts.jsonl",
}


def read_jsonl(path: Path) -> list[dict]:
    if not path.exists():
        return []
    rows: list[dict] = []
    for line in path.read_text(encoding="utf-8-sig", errors="replace").splitlines():
        if not line.strip():
            continue
        value = json.loads(line)
        if isinstance(value, dict):
            rows.append(value)
    return rows


def latest_submission(root: Path) -> Path:
    candidates = sorted(
        (path for path in root.glob("*_zhihu01") if path.is_dir()),
        key=lambda path: path.name,
    )
    if not candidates:
        raise FileNotFoundError(f"No Zhihu submission directory under {root}")
    return candidates[-1]


def yes_no(value) -> str:
    if value in (True, "true", "True", "是", 1, "1"):
        return "是"
    if value in (False, "false", "False", "否", 0, "0"):
        return "否"
    return ""


def blank_if_none(value):
    return "" if value is None else value


def write_rows(ws, rows: list[dict], fields: list[str], transforms: dict[str, object] | None = None) -> None:
    transforms = transforms or {}
    for row_index, row in enumerate(rows, start=2):
        for column_index, field in enumerate(fields, start=1):
            value = row.get(field, "")
            transform = transforms.get(field)
            if callable(transform):
                value = transform(value)
            cell = ws.cell(row=row_index, column=column_index, value=blank_if_none(value))
            if field.endswith("_id") or field in {"content_id", "comment_id"}:
                cell.number_format = "@"


def export_workbook(template: Path, submission: Path, output: Path) -> dict:
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template, output)
    workbook = load_workbook(output)

    tables = {
        sheet: read_jsonl(submission / filename)
        for sheet, filename in TABLE_FILES.items()
    }

    write_rows(
        workbook["表1-发布内容"],
        tables["表1-发布内容"],
        [
            "content_id", "platform", "author_id", "author_name",
            "author_ip_location", "account_type", "content_type", "title",
            "content_text", "publish_time", "collection_time", "original_url",
            "matched_keywords", "original_or_repost", "is_valid_monitoring_data",
            "invalid_reason",
        ],
        {
            "matched_keywords": lambda value: "；".join(value) if isinstance(value, list) else value,
            "is_valid_monitoring_data": yes_no,
        },
    )
    write_rows(
        workbook["表2-评论基础"],
        tables["表2-评论基础"],
        [
            "content_id", "comment_id", "platform", "comment_user_id",
            "comment_user_ip_location", "comment_text", "comment_publish_time",
            "is_valid_comment",
        ],
        {"is_valid_comment": yes_no},
    )
    write_rows(
        workbook["表3-发布互动"],
        tables["表3-发布互动"],
        [
            "content_id", "platform", "snapshot_time", "view_count", "like_count",
            "comment_count", "repost_count", "share_count", "favorite_count",
        ],
    )
    write_rows(
        workbook["表4-评论互动"],
        tables["表4-评论互动"],
        [
            "content_id", "comment_id", "platform", "snapshot_time",
            "comment_reply_count", "comment_like_count",
        ],
    )

    account_sheet = workbook["表5-账号信息"]
    existing_by_name = {
        str(account_sheet.cell(row=i, column=3).value or "").strip(): i
        for i in range(2, account_sheet.max_row + 1)
        if str(account_sheet.cell(row=i, column=3).value or "").strip()
    }
    next_row = max(existing_by_name.values(), default=1) + 1
    account_fields = [
        "author_id", "platform", "author_name", "profile_url", "account_type",
        "fan_count", "following_count", "region", "organization", "is_key_account",
        "related_post_count", "view_count", "like_count", "comment_count",
        "repost_count", "favorite_count", "total_engagement", "collection_time",
    ]
    for account in tables["表5-账号信息"]:
        name = str(account.get("author_name") or "").strip()
        row_index = existing_by_name.get(name)
        if row_index is None:
            row_index = next_row
            next_row += 1
        for column_index, field in enumerate(account_fields, start=1):
            value = account.get(field, "")
            if field == "is_key_account":
                value = yes_no(value)
            cell = account_sheet.cell(row=row_index, column=column_index, value=blank_if_none(value))
            if field == "author_id":
                cell.number_format = "@"

    try:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"
    except AttributeError:
        pass
    workbook.save(output)

    row_counts = {sheet: len(rows) for sheet, rows in tables.items()}
    manifest = {
        "template": str(template),
        "source_submission": str(submission),
        "output_workbook": str(output),
        "row_counts": row_counts,
    }
    manifest_path = output.with_suffix(".manifest.json")
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return manifest


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument(
        "--submission-root",
        type=Path,
        default=REPO_ROOT / "data_submissions" / "zhihu",
    )
    parser.add_argument("--submission", type=Path)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    submission = args.submission or latest_submission(args.submission_root)
    result = export_workbook(args.template.resolve(), submission.resolve(), args.output.resolve())
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
