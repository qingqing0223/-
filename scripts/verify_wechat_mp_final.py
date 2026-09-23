"""Configuration / exported data checks, without importing the old classifier."""
from __future__ import annotations

import argparse
import csv
from datetime import datetime
import json
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from wechat.mp_export import FIELDS, FILES, SHEETS, WORKBOOK
from wechat.mp_poms import POMS_FIELDS
from wechat.mp_records import START, assess_record, time_problem
from wechat.mp_display import empty, normalize_tables
from wechat.mp_poms import validate_schema

EXPECTED_KEYWORDS = ["2026年民族团结进步宣传周", "首个民族团结进步宣传周", "促进民族团结进步，奋进伟大复兴征程", "民族团结进步倡议", "民族团结进步宣传周主场活动", "石榴花开——铸牢中华民族共同体意识"]


def verify_export(directory: Path) -> dict:
    from openpyxl import load_workbook  # Read-only independent verification.
    rows = [json.loads(line) for line in (directory / "search_contents.jsonl").read_text(encoding="utf-8").splitlines() if line]
    assert len({r["content_id"] for r in rows}) == len(rows), "duplicate IDs"
    assert all(r.get("matched_keywords") for r in rows), "missing keywords"
    summary = json.loads((directory / "summary.json").read_text(encoding="utf-8"))
    book = load_workbook(directory / summary.get("workbook_file", WORKBOOK), data_only=False)
    assert book.sheetnames == ["说明与统计", *SHEETS], book.sheetnames
    normalized = summary.get("display_normalization_version") == 2
    all_batches = []
    assert all(not time_problem(r, summary["monitoring_start_time"], summary.get("monitoring_end_time")) for r in rows), "out-of-time monitoring candidate"
    for cell, key in (("B2", "monitoring_start_time"), ("B5", "content_export_at"), ("B6", "metrics_export_at")):
        expected = datetime.fromisoformat(summary[key]).replace(tzinfo=None)
        assert abs((book["说明与统计"][cell].value - expected).total_seconds()) < 0.01
    counts = []
    for index, (filename, fields, name) in enumerate(zip(FILES, FIELDS, SHEETS)):
        with (directory / filename).open(encoding="utf-8-sig", newline="") as stream:
            data = list(csv.reader(stream))
        assert data[0] == fields, filename
        assert all(len(r) == len(fields) for r in data), filename
        sheet = book[name]
        assert [sheet.cell(1, c + 1).value for c in range(len(fields))] == fields, name
        populated = list(sheet.iter_rows(min_row=2, max_col=len(fields))) if sheet.max_row > 1 else []
        assert len(populated) == len(data) - 1, name
        for cells, csv_row in zip(populated, data[1:]):
            for col, cell in enumerate(cells):
                assert cell.data_type != "f", (name, cell.coordinate)
                if normalized:
                    assert not empty(cell.value), (name, cell.coordinate, "empty data cell")
                if "编号" in fields[col] or "ID" in fields[col]:
                    assert cell.number_format == "@", (name, cell.coordinate)
                    assert cell.value is None or isinstance(cell.value, str)
                if csv_row[col] == "":
                    assert cell.value is None, (name, cell.coordinate, cell.value)
                elif "时间" in fields[col]:
                    expected = datetime.fromisoformat(csv_row[col]).replace(tzinfo=None)
                    assert isinstance(cell.value, datetime) and abs((cell.value - expected).total_seconds()) < 0.01, (name, cell.coordinate)
                else:
                    actual = str(cell.value)
                    # Artifact Tool may preserve the explicit literal escape.
                    assert actual == csv_row[col] or "'" + actual == csv_row[col] or actual == "'" + csv_row[col], (name, cell.coordinate, actual, csv_row[col])
        if index in (1, 3):
            assert len(data) == 1, "Public source should not invent comments"
        counts.append(len(data) - 1)
        if index == 0:
            for csv_row, original in zip(data[1:], rows):
                assert csv_row[0] == original["content_id"]
                assert csv_row[3] == (original.get("author") or "")
                assert csv_row[9] == original["publish_time"]
                assert csv_row[10] == original["collected_at"]
                assert json.loads(csv_row[12]) == original["matched_keywords"]
                assessed = assess_record(original, summary["monitoring_start_time"], end=summary.get("monitoring_end_time"))
                assert csv_row[14] == assessed["review_status"]
                assert csv_row[15] == (assessed["invalid_reason"] or ("无" if normalized else ""))
        batches = json.loads((directory / f"table{index+1}_batch.json").read_text(encoding="utf-8"))
        assert isinstance(batches, list) and len(batches) == len(data) - 1
        all_batches.append(batches)
        for obj, cells in zip(batches, populated):
            assert list(obj) == POMS_FIELDS[index]
            for field, cell in zip(POMS_FIELDS[index], cells):
                if field.endswith("_id"):
                    assert obj[field] is None or isinstance(obj[field], str)
                if cell.value is None:
                    assert obj[field] is None or obj[field] == ""
                elif field == "matched_keywords":
                    assert obj[field] == json.loads(cell.value)
                elif field.startswith("is_"):
                    assert obj[field] is {"是": True, "否": False, "待核验": False if normalized else None}[cell.value]
                elif field.endswith("_count"):
                    assert type(obj[field]) is int and obj[field] == cell.value
    book.close()
    if normalized:
        validate_schema(all_batches)
        account_ids = {r["account_name"]: r["account_id"] for r in all_batches[4]}
        assert all(r["publisher_account_id"] == account_ids[r["account_name"]] for r in all_batches[0] if r["account_name"] in account_ids)
    assert counts[0] == len(rows), "JSONL / CSV row mismatch"
    assert summary["total_candidates"] == len(rows)
    for label, field in (("是", "valid_articles"), ("否", "invalid_articles"), ("待核验", "pending_review_articles")):
        assert summary[field] == sum(r["review_status"] == label for r in rows)
    return {"ok": True, "table_counts": counts, "jsonl_records": len(rows), "text_ids_and_nulls_checked": True}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", default=str(ROOT / "config/monitoring.wechat.windows.json"))
    parser.add_argument("--config-only", action="store_true")
    parser.add_argument("--directory", type=Path)
    args = parser.parse_args()
    cfg = json.loads(Path(args.config).read_text(encoding="utf-8-sig"))
    from wechat.mp_pipeline import validate_config
    validate_config(cfg)
    assert cfg["wechat_mp_search_until_exhausted"]
    assert int(cfg.get("wechat_mp_content_export_seconds", 900)) > 0
    assert int(cfg.get("wechat_mp_interaction_export_seconds", 3600)) > 0
    catalog = Path(cfg.get("wechat_mp_key_accounts", "config/key_accounts.wechat_mp.json"))
    assert catalog.is_file(), str(catalog)
    result = {"ok": True, "mode": "config-only"}
    if not args.config_only:
        directory = args.directory
        if directory is None:
            candidates = sorted(Path(cfg.get("wechat_mp_submission_root", "data_submissions/wechat_mp")).glob("*_*"))
            directory = candidates[-1] if candidates else None
        assert directory and directory.is_dir(), "No V3 submission directory"
        result = verify_export(directory)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
