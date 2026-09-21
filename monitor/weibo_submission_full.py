from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from pipeline.normalizer import normalize_record
from monitor.weibo_submission import (
    _load_account_types,
    _read_rows,
    _write_rows,
    export_weibo_table1,
)


TABLE_FILES = {
    "table1": "table1_content.jsonl",
    "table2": "table2_comments.jsonl",
    "table3": "table3_post_interactions.jsonl",
    "table4": "table4_comment_interactions.jsonl",
    "table5": "table5_accounts.jsonl",
}


def _clean(value) -> str:
    return "" if value is None else str(value).strip()


def _number(value):
    if value in (None, ""):
        return ""

    if isinstance(value, bool):
        return ""

    if isinstance(value, (int, float)):
        return int(value)

    text = str(value).strip().replace(",", "")

    if not text:
        return ""

    multiplier = 1

    if text.endswith("\u4e07"):
        multiplier = 10000
        text = text[:-1]
    elif text.lower().endswith("w"):
        multiplier = 10000
        text = text[:-1]
    elif text.lower().endswith("k"):
        multiplier = 1000
        text = text[:-1]

    try:
        return int(float(text) * multiplier)
    except (TypeError, ValueError):
        return ""


def _public(raw: dict, *names):
    for name in names:
        value = raw.get(name)
        if value not in (None, ""):
            return value
    return ""


def _parse_time(value):
    text = _clean(value).replace("Z", "+00:00")

    if not text:
        return None

    try:
        return datetime.fromisoformat(text)
    except ValueError:
        return None


def _hour(value) -> str:
    dt = value if isinstance(value, datetime) else _parse_time(value)

    if dt is None:
        return _clean(value)

    return dt.replace(
        minute=0,
        second=0,
        microsecond=0,
    ).isoformat(timespec="seconds")


def _safe_node(value: str) -> str:
    out = "".join(
        ch if ch.isalnum() or ch in "._-"
        else "_"
        for ch in _clean(value)
    )
    return out or "wb01"


def _all_raw_files(
    files: list[Path],
    raw_root: Path | None,
) -> list[Path]:
    found = {
        Path(path).resolve()
        for path in files
        if Path(path).exists()
    }

    if raw_root and raw_root.exists():
        for path in raw_root.rglob("*.jsonl"):
            found.add(path.resolve())

    return sorted(found)


def _normalized_rows(files: list[Path]):
    for path in files:
        if "account_info" in path.name.lower():
            continue

        for raw in _read_rows(path):
            rec = normalize_record(
                raw,
                source_file=path.name,
                platform_hint="wb",
            )

            if not rec:
                continue

            if rec.get("platform") != "wb":
                continue

            yield raw, rec


def _upsert(
    path: Path,
    rows: list[dict],
    key_fields: tuple[str, ...],
) -> None:
    existing = {}

    for row in _read_rows(path):
        key = tuple(
            _clean(row.get(field))
            for field in key_fields
        )

        if all(key):
            existing[key] = row

    for row in rows:
        key = tuple(
            _clean(row.get(field))
            for field in key_fields
        )

        if not all(key):
            continue

        old = existing.get(key, {})
        merged = dict(old)

        for field, value in row.items():
            if field not in merged or value not in (None, "", []):
                merged[field] = value

        existing[key] = merged

    _write_rows(
        path,
        [
            existing[key]
            for key in sorted(existing)
        ],
    )


def _replace_hourly(
    path: Path,
    incoming: list[dict],
    id_fields: tuple[str, ...],
    valid_ids: set[str],
    valid_field: str,
) -> None:
    existing = {}

    for row in _read_rows(path):
        if _clean(row.get(valid_field)) not in valid_ids:
            continue

        key = tuple(
            [_clean(row.get(x)) for x in id_fields]
            + [_hour(row.get("snapshot_time"))]
        )

        existing[key] = row

    for row in incoming:
        if _clean(row.get(valid_field)) not in valid_ids:
            continue

        key = tuple(
            [_clean(row.get(x)) for x in id_fields]
            + [_hour(row.get("snapshot_time"))]
        )

        existing[key] = row

    _write_rows(
        path,
        [
            existing[key]
            for key in sorted(existing)
        ],
    )


def _profiles(files: list[Path]):
    by_id = {}
    by_name = {}

    for path in files:
        if "account_info" not in path.name.lower():
            continue

        for row in _read_rows(path):
            account_id = _clean(
                row.get("account_id")
                or row.get("creator_id")
            )

            account_name = _clean(
                row.get("account_name")
                or row.get("name")
            )

            if account_id:
                by_id[account_id] = row

            if account_name:
                by_name[account_name] = row

    return by_id, by_name



def _looks_masked(value) -> bool:
    return "*" in _clean(value)


def _seed_rows(
    seed_dir: Path | None,
    table_name: str,
) -> list[dict]:
    if seed_dir is None:
        return []

    filename = TABLE_FILES[table_name]
    return _read_rows(seed_dir / filename)


def _merge_seed_table1(
    current: list[dict],
    seed: list[dict],
) -> list[dict]:
    by_id = {
        _clean(row.get("content_id")): dict(row)
        for row in current
        if _clean(row.get("content_id"))
    }

    for source in seed:
        content_id = _clean(
            source.get("content_id")
        )

        if not content_id:
            continue

        if content_id not in by_id:
            by_id[content_id] = dict(source)
            continue

        row = by_id[content_id]

        current_id = _clean(
            row.get("author_id")
        )
        seed_id = _clean(
            source.get("author_id")
        )

        # Prefer a platform-public numeric Weibo ID over a legacy hash.
        if (
            seed_id.isdigit()
            and not current_id.isdigit()
        ):
            row["author_id"] = seed_id

        current_name = _clean(
            row.get("author_name")
        )
        seed_name = _clean(
            source.get("author_name")
        )

        # Prefer an already verified full public display name.
        if (
            seed_name
            and not _looks_masked(seed_name)
            and (
                not current_name
                or _looks_masked(current_name)
            )
        ):
            row["author_name"] = seed_name

        for field in (
            "author_ip_location",
            "account_type",
            "content_type",
            "title",
            "content_text",
            "publish_time",
            "original_url",
            "original_or_repost",
        ):
            if (
                row.get(field) in (None, "", [])
                and source.get(field)
                not in (None, "", [])
            ):
                row[field] = source[field]

        matches = []
        seen = set()

        for value in (
            list(row.get("matched_keywords") or [])
            + list(source.get("matched_keywords") or [])
        ):
            value = _clean(value)

            if value and value not in seen:
                seen.add(value)
                matches.append(value)

        row["matched_keywords"] = matches

        # The seed consists of the already manually accepted historical set.
        if source.get(
            "is_valid_monitoring_data"
        ) is True:
            row["is_valid_monitoring_data"] = True
            row["invalid_reason"] = ""

        by_id[content_id] = row

    return [
        by_id[key]
        for key in sorted(by_id)
    ]


def _merge_seed_table2(
    current: list[dict],
    seed: list[dict],
) -> list[dict]:
    by_id = {
        _clean(row.get("comment_id")): dict(row)
        for row in current
        if _clean(row.get("comment_id"))
    }

    for source in seed:
        comment_id = _clean(
            source.get("comment_id")
        )

        if not comment_id:
            continue

        if comment_id not in by_id:
            by_id[comment_id] = dict(source)
            continue

        row = by_id[comment_id]

        current_id = _clean(
            row.get("comment_user_id")
        )
        seed_id = _clean(
            source.get("comment_user_id")
        )

        if (
            seed_id.isdigit()
            and not current_id.isdigit()
        ):
            row["comment_user_id"] = seed_id

        current_name = _clean(
            row.get("comment_user_name")
        )
        seed_name = _clean(
            source.get("comment_user_name")
        )

        if (
            seed_name
            and not _looks_masked(seed_name)
            and (
                not current_name
                or _looks_masked(current_name)
            )
        ):
            row["comment_user_name"] = seed_name

        for field in (
            "content_id",
            "platform",
            "comment_user_ip_location",
            "comment_text",
            "comment_publish_time",
            "parent_comment_id",
            "root_comment_id",
            "comment_level",
            "sub_comment_count",
            "comment_like_count",
        ):
            if (
                row.get(field) in (None, "", [])
                and source.get(field)
                not in (None, "", [])
            ):
                row[field] = source[field]

        if source.get(
            "is_valid_comment"
        ) is True:
            row["is_valid_comment"] = True

        by_id[comment_id] = row

    return [
        by_id[key]
        for key in sorted(by_id)
    ]

def _build_table5(
    table1: list[dict],
    table3: list[dict],
    previous: list[dict],
    files: list[Path],
    cfg: dict,
    repo_root: Path,
    now_text: str,
) -> list[dict]:
    account_types = _load_account_types(
        repo_root,
        cfg,
    )

    profiles_by_id, profiles_by_name = _profiles(files)

    old_by_id = {}
    old_by_name = {}

    for source in previous:
        account_id = _clean(
            source.get("account_id")
        )
        account_name = _clean(
            source.get("account_name")
        )

        targets = []

        if account_id:
            targets.append(
                (old_by_id, account_id)
            )

        if account_name:
            targets.append(
                (old_by_name, account_name)
            )

        for mapping, key in targets:
            merged = dict(
                mapping.get(key, {})
            )

            for field, value in source.items():
                # A later empty historical/current row must never erase
                # already collected public profile information.
                if (
                    field not in merged
                    or value not in (None, "", [])
                ):
                    merged[field] = value

            mapping[key] = merged

    latest_engagement = {}

    for row in table3:
        cid = _clean(row.get("content_id"))

        if not cid:
            continue

        current = latest_engagement.get(cid)

        if (
            current is None
            or _clean(row.get("snapshot_time"))
            >= _clean(current.get("snapshot_time"))
        ):
            latest_engagement[cid] = row

    accounts = {}

    for post in table1:
        if post.get("is_valid_monitoring_data") is not True:
            continue

        account_id = _clean(post.get("author_id"))
        account_name = _clean(post.get("author_name"))

        if not account_id and not account_name:
            continue

        key = (
            "id:" + account_id
            if account_id
            else "name:" + account_name
        )

        item = accounts.setdefault(
            key,
            {
                "account_id": account_id,
                "account_name": account_name,
                "content_ids": set(),
            },
        )

        if account_name:
            item["account_name"] = account_name

        content_id = _clean(post.get("content_id"))

        if content_id:
            item["content_ids"].add(content_id)

    output = []

    for item in accounts.values():
        account_id = item["account_id"]
        account_name = item["account_name"]

        profile = (
            profiles_by_id.get(account_id)
            or profiles_by_name.get(account_name)
            or {}
        )

        old = (
            old_by_id.get(account_id)
            or old_by_name.get(account_name)
            or {}
        )

        profile_url = (
            _clean(profile.get("profile_url"))
            or _clean(old.get("profile_url"))
        )

        if account_id.isdigit():
            profile_url = "https://m.weibo.cn/u/" + account_id

        followers = _number(profile.get("followers"))
        following = _number(profile.get("following"))

        if followers == "":
            followers = old.get("followers", "")

        if following == "":
            following = old.get("following", "")

        account_type = (
            account_types.get(account_name, "")
            or _clean(old.get("account_type"))
        )

        region = (
            _clean(profile.get("region"))
            or _clean(old.get("region"))
        )

        organization = (
            _clean(profile.get("organization"))
            or _clean(profile.get("auth_info"))
            or _clean(old.get("organization"))
        )

        likes = 0
        comments = 0
        reposts = 0

        for content_id in item["content_ids"]:
            row = latest_engagement.get(content_id, {})

            likes += int(row.get("like_count") or 0)
            comments += int(row.get("comment_count") or 0)
            reposts += int(row.get("repost_count") or 0)

        output.append({
            "account_id": account_id,
            "platform": "\u5fae\u535a",
            "account_name": account_name,
            "profile_url": profile_url,
            "account_type": account_type,
            "followers": followers,
            "following": following,
            "region": region,
            "organization": organization,
            "is_key_account": bool(
                account_types.get(account_name)
            ),
            "related_post_count": len(item["content_ids"]),
            "views": "",
            "likes": likes,
            "comments": comments,
            "shares": reposts,
            "favorites": "",
            "total_interactions": likes + comments + reposts,
            "collected_at": now_text,
        })

    return sorted(
        output,
        key=lambda x: (
            _clean(x.get("account_name")),
            _clean(x.get("account_id")),
        ),
    )


def _excel(
    path: Path,
    manifest: dict,
    tables: dict[str, list[dict]],
) -> None:
    wb = Workbook()

    ws = wb.active
    ws.title = "\u8bf4\u660e\u4e0e\u7edf\u8ba1"

    ws.append(["\u9879\u76ee", "\u503c"])
    ws.append(["\u5e73\u53f0", "\u5fae\u535a"])
    ws.append(["\u751f\u6210\u65f6\u95f4", manifest["generated_at"]])
    ws.append(["\u88681\u6570\u91cf", manifest["table1_rows"]])
    ws.append(["\u88682\u6570\u91cf", manifest["table2_rows"]])
    ws.append(["\u88683\u6570\u91cf", manifest["table3_rows"]])
    ws.append(["\u88684\u6570\u91cf", manifest["table4_rows"]])
    ws.append(["\u88685\u6570\u91cf", manifest["table5_rows"]])
    ws.append([
        "\u8bf4\u660e",
        "\u5e73\u53f0\u672a\u516c\u5f00\u6216\u65e0\u6cd5\u53ef\u9760\u83b7\u53d6\u7684\u5b57\u6bb5\u7559\u7a7a\uff0c\u4e0d\u4f7f\u7528\u4f30\u7b97\u503c\u6216\u66ff\u4ee30\u3002",
    ])

    specs = {
        "\u88681-\u53d1\u5e03\u5185\u5bb9": (
            "table1",
            [
                ("content_id", "\u53d1\u5e03\u5185\u5bb9\u7f16\u53f7"),
                ("platform", "\u5e73\u53f0"),
                ("author_id", "\u53d1\u5e03\u5e10\u53f7ID"),
                ("author_name", "\u5e10\u53f7\u540d\u79f0"),
                ("author_ip_location", "\u5e10\u53f7IP\u5c5e\u5730"),
                ("account_type", "\u5e10\u53f7\u7c7b\u578b"),
                ("content_type", "\u5185\u5bb9\u7c7b\u578b"),
                ("title", "\u6807\u9898"),
                ("content_text", "\u6b63\u6587/\u6587\u6848"),
                ("publish_time", "\u53d1\u5e03\u65f6\u95f4"),
                ("collection_time", "\u91c7\u96c6\u65f6\u95f4"),
                ("original_url", "\u539f\u59cb\u94fe\u63a5"),
                ("matched_keywords", "\u547d\u4e2d\u5173\u952e\u8bcd"),
                ("original_or_repost", "\u539f\u521b/\u8f6c\u8f7d"),
                ("is_valid_monitoring_data", "\u662f\u5426\u6709\u6548"),
                ("invalid_reason", "\u65e0\u6548\u539f\u56e0"),
            ],
        ),
        "\u88682-\u8bc4\u8bba": (
            "table2",
            [
                ("content_id", "\u53d1\u5e03\u5185\u5bb9\u7f16\u53f7"),
                ("comment_id", "\u8bc4\u8bba\u7f16\u53f7"),
                ("platform", "\u5e73\u53f0"),
                ("comment_user_id", "\u8bc4\u8bba\u7528\u6237ID"),
                ("comment_user_name", "\u8bc4\u8bba\u7528\u6237\u540d"),
                ("comment_user_ip_location", "\u8bc4\u8bba\u7528\u6237IP\u5c5e\u5730"),
                ("comment_text", "\u8bc4\u8bba\u6b63\u6587"),
                ("comment_publish_time", "\u8bc4\u8bba\u53d1\u5e03\u65f6\u95f4"),
                ("is_valid_comment", "\u662f\u5426\u6709\u6548"),
                ("parent_comment_id", "\u7236\u8bc4\u8bbaID"),
                ("root_comment_id", "\u6839\u8bc4\u8bbaID"),
                ("comment_level", "\u8bc4\u8bba\u5c42\u7ea7"),
                ("sub_comment_count", "\u56de\u590d\u6570"),
                ("comment_like_count", "\u8bc4\u8bba\u70b9\u8d5e\u6570"),
            ],
        ),
        "\u88683-\u53d1\u5e03\u4e92\u52a8": (
            "table3",
            [
                ("content_id", "\u53d1\u5e03\u5185\u5bb9\u7f16\u53f7"),
                ("platform", "\u5e73\u53f0"),
                ("snapshot_time", "\u7edf\u8ba1\u65f6\u95f4"),
                ("view_count", "\u9605\u8bfb/\u64ad\u653e\u91cf"),
                ("like_count", "\u70b9\u8d5e\u91cf"),
                ("comment_count", "\u8bc4\u8bba\u91cf"),
                ("repost_count", "\u8f6c\u53d1\u91cf"),
                ("share_count", "\u5206\u4eab\u91cf"),
                ("favorite_count", "\u6536\u85cf\u91cf"),
            ],
        ),
        "\u88684-\u8bc4\u8bba\u4e92\u52a8": (
            "table4",
            [
                ("content_id", "\u53d1\u5e03\u5185\u5bb9\u7f16\u53f7"),
                ("comment_id", "\u8bc4\u8bba\u7f16\u53f7"),
                ("platform", "\u5e73\u53f0"),
                ("snapshot_time", "\u7edf\u8ba1\u65f6\u95f4"),
                ("comment_reply_count", "\u8bc4\u8bba\u56de\u590d\u6570"),
                ("comment_like_count", "\u8bc4\u8bba\u70b9\u8d5e\u6570"),
            ],
        ),
        "\u88685-\u5e10\u53f7\u4fe1\u606f": (
            "table5",
            [
                ("account_id", "\u5e10\u53f7ID"),
                ("platform", "\u5e73\u53f0"),
                ("account_name", "\u5e10\u53f7\u540d\u79f0"),
                ("profile_url", "\u4e3b\u9875\u5730\u5740"),
                ("account_type", "\u5e10\u53f7\u7c7b\u578b"),
                ("followers", "\u7c89\u4e1d\u91cf"),
                ("following", "\u5173\u6ce8\u91cf"),
                ("region", "\u6240\u5c5e\u5730\u533a"),
                ("organization", "\u6240\u5c5e\u673a\u6784"),
                ("is_key_account", "\u91cd\u70b9\u76d1\u6d4b\u5e10\u53f7"),
                ("related_post_count", "\u76f8\u5173\u53d1\u6587\u91cf"),
                ("views", "\u9605\u8bfb/\u64ad\u653e\u91cf"),
                ("likes", "\u70b9\u8d5e\u91cf"),
                ("comments", "\u8bc4\u8bba\u91cf"),
                ("shares", "\u8f6c\u53d1\u91cf"),
                ("favorites", "\u6536\u85cf\u91cf"),
                ("total_interactions", "\u603b\u4e92\u52a8\u91cf"),
                ("collected_at", "\u91c7\u96c6\u65f6\u95f4"),
            ],
        ),
    }

    for sheet_name, (table_name, columns) in specs.items():
        sheet = wb.create_sheet(sheet_name)

        sheet.append([header for _, header in columns])

        for cell in sheet[1]:
            cell.font = Font(bold=True)

        for row in tables[table_name]:
            values = []

            for key, _ in columns:
                value = row.get(key, "")

                if isinstance(value, list):
                    value = "\uff1b".join(map(str, value))

                if isinstance(value, bool):
                    value = "\u662f" if value else "\u5426"

                values.append(value)

            sheet.append(values)

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        for index, (_, header) in enumerate(columns, start=1):
            letter = get_column_letter(index)

            for cell in sheet[letter]:
                if "ID" in header or "\u7f16\u53f7" in header:
                    cell.number_format = "@"

            sheet.column_dimensions[letter].width = min(
                42,
                max(12, len(header) + 4),
            )

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def export_weibo_submission(
    files: list[Path],
    cfg: dict,
    repo_root: Path,
    node_id: str = "wb01",
    raw_root: Path | None = None,
    seed_dir: Path | None = None,
    now: datetime | None = None,
) -> dict:
    now_dt = now or datetime.now().astimezone()

    if now_dt.tzinfo is None:
        now_dt = now_dt.astimezone()

    now_text = now_dt.isoformat(timespec="seconds")
    node = _safe_node(node_id)

    output_dir = (
        repo_root
        / "data_submissions"
        / "weibo"
        / f"{now_dt.date().isoformat()}_{node}"
    )

    output_dir.mkdir(parents=True, exist_ok=True)

    paths = {
        key: output_dir / name
        for key, name in TABLE_FILES.items()
    }

    all_files = _all_raw_files(files, raw_root)

    if seed_dir is not None:
        seed_dir = Path(seed_dir)

    seed_table1 = _seed_rows(seed_dir, "table1")
    seed_table2 = _seed_rows(seed_dir, "table2")
    seed_table3 = _seed_rows(seed_dir, "table3")
    seed_table4 = _seed_rows(seed_dir, "table4")
    seed_table5 = _seed_rows(seed_dir, "table5")

    export_weibo_table1(
        all_files,
        paths["table1"],
        cfg,
        repo_root,
    )

    table1 = _merge_seed_table1(
        _read_rows(paths["table1"]),
        seed_table1,
    )

    # Formal submission contains accepted monitoring rows only.
    table1 = [
        row
        for row in table1
        if row.get("is_valid_monitoring_data") is True
    ]

    _write_rows(
        paths["table1"],
        table1,
    )

    valid_content_ids = {
        _clean(row.get("content_id"))
        for row in table1
        if row.get("is_valid_monitoring_data") is True
    }

    table2_incoming = []
    table3_incoming = []
    table4_incoming = []

    for raw, rec in _normalized_rows(all_files):
        if rec.get("record_type") == "comment":
            content_id = _clean(rec.get("content_id"))
            comment_id = _clean(
                rec.get("comment_id")
                or rec.get("sample_id")
            )

            if (
                not content_id
                or content_id not in valid_content_ids
                or not comment_id
            ):
                continue

            table2_incoming.append({
                "content_id": content_id,
                "comment_id": comment_id,
                "platform": "\u5fae\u535a",
                "comment_user_id": _clean(rec.get("author_id")),
                "comment_user_name": _clean(rec.get("author")),
                "comment_user_ip_location": _clean(rec.get("ip_location")),
                "comment_text": _clean(rec.get("content")),
                "comment_publish_time": _clean(rec.get("publish_time")),
                "is_valid_comment": True,
                "parent_comment_id": _clean(rec.get("parent_comment_id")),
                "root_comment_id": _clean(rec.get("root_comment_id")),
                "comment_level": int(rec.get("comment_level") or 1),
                "sub_comment_count": _number(
                    _public(
                        raw,
                        "sub_comment_count",
                        "total_number",
                        "reply_count",
                    )
                ),
                "comment_like_count": _number(
                    _public(
                        raw,
                        "comment_like_count",
                        "like_count",
                        "liked_count",
                    )
                ),
            })

            table4_incoming.append({
                "content_id": content_id,
                "comment_id": comment_id,
                "platform": "\u5fae\u535a",
                "snapshot_time": now_text,
                "comment_reply_count": _number(
                    _public(
                        raw,
                        "sub_comment_count",
                        "total_number",
                        "reply_count",
                    )
                ),
                "comment_like_count": _number(
                    _public(
                        raw,
                        "comment_like_count",
                        "like_count",
                        "liked_count",
                    )
                ),
            })

            continue

        content_id = _clean(
            rec.get("content_id")
            or rec.get("sample_id")
        )

        if content_id not in valid_content_ids:
            continue

        table3_incoming.append({
            "content_id": content_id,
            "platform": "\u5fae\u535a",
            "snapshot_time": now_text,

            # Missing public metrics remain blank.
            "view_count": _number(
                _public(
                    raw,
                    "view_count",
                    "views",
                    "play_count",
                )
            ),

            "like_count": _number(
                _public(
                    raw,
                    "liked_count",
                    "attitudes_count",
                    "like_count",
                    "likes",
                )
            ),

            "comment_count": _number(
                _public(
                    raw,
                    "comments_count",
                    "comment_count",
                    "comments",
                )
            ),

            # MediaCrawler shared_count is Weibo repost count.
            "repost_count": _number(
                _public(
                    raw,
                    "shared_count",
                    "reposts_count",
                    "repost_count",
                    "forward_count",
                )
            ),

            # Never duplicate shared_count into independent share_count.
            "share_count": _number(
                _public(
                    raw,
                    "share_count",
                    "shares",
                )
            ),

            "favorite_count": _number(
                _public(
                    raw,
                    "favorite_count",
                    "favorites",
                    "collected_count",
                )
            ),
        })

    existing_table2 = [
        row
        for row in _read_rows(paths["table2"])
        if _clean(row.get("content_id")) in valid_content_ids
    ]

    _write_rows(paths["table2"], existing_table2)

    _upsert(
        paths["table2"],
        table2_incoming,
        ("comment_id",),
    )

    table2 = _merge_seed_table2(
        _read_rows(paths["table2"]),
        seed_table2,
    )

    table2 = [
        row
        for row in table2
        if (
            row.get("is_valid_comment") is True
            and _clean(row.get("content_id"))
            in valid_content_ids
        )
    ]

    _write_rows(
        paths["table2"],
        table2,
    )

    valid_comment_ids = {
        _clean(row.get("comment_id"))
        for row in table2
        if row.get("is_valid_comment") is True
    }

    _replace_hourly(
        paths["table3"],
        seed_table3 + table3_incoming,
        ("content_id",),
        valid_content_ids,
        "content_id",
    )

    _replace_hourly(
        paths["table4"],
        seed_table4 + table4_incoming,
        ("content_id", "comment_id"),
        valid_comment_ids,
        "comment_id",
    )

    table3 = _read_rows(paths["table3"])
    table4 = _read_rows(paths["table4"])

    table5 = _build_table5(
        table1,
        table3,
        seed_table5 + _read_rows(paths["table5"]),
        all_files,
        cfg,
        repo_root,
        now_text,
    )

    _write_rows(paths["table5"], table5)

    tables = {
        "table1": table1,
        "table2": table2,
        "table3": table3,
        "table4": table4,
        "table5": table5,
    }

    manifest = {
        "platform": "wb",
        "node_id": node,
        "generated_at": now_text,
        "monitoring_start_time": _clean(
            cfg.get("monitoring_start_time")
        ),
        "classification": "not_run_collection_group_scope_only",
        "table1_rows": len(table1),
        "table2_rows": len(table2),
        "table3_rows": len(table3),
        "table4_rows": len(table4),
        "table5_rows": len(table5),
        "snapshot_hour": _hour(now_dt),
        "unsupported_public_fields_note": (
            "Weibo metrics unavailable from the public collection "
            "interface remain blank; no estimates or substitute zeros."
        ),
    }

    manifest_path = output_dir / "manifest.json"

    manifest_path.write_text(
        json.dumps(
            manifest,
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )

    excel_path = (
        output_dir
        / "\u5fae\u535a\u76d1\u6d4b\u6570\u636e_TechDesignV3.xlsx"
    )

    _excel(
        excel_path,
        manifest,
        tables,
    )

    return {
        **manifest,
        "accepted_rows": len(table1) + len(table2),
        "output_dir": str(output_dir),
        "excel": str(excel_path),
        "manifest": str(manifest_path),
        "table_files": {
            key: str(value)
            for key, value in paths.items()
        },
    }